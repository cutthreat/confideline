from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "etalon-tz-g2-geo-content-restoration.md"
HTML_OUT = ROOT / "tz-g2-geo-content-for-igor.html"
DOCX_OUT = ROOT / "tz-g2-geo-content-for-igor.docx"


def parse_blocks(text: str):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            yield ("code", block)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            yield ("heading", (len(heading.group(1)), heading.group(2)))
            i += 1
            continue
        if line.lstrip().startswith("|"):
            table = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{3,}", lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    table.append(cells)
                i += 1
            if table:
                yield ("table", table)
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        number = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if bullet:
            yield ("bullet", bullet.group(1))
            i += 1
            continue
        if number:
            yield ("number", (int(number.group(1)), number.group(2)))
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip() or nxt.startswith("#") or nxt.startswith("```") or nxt.lstrip().startswith("|") or re.match(r"^\s*[-*]\s+", nxt) or re.match(r"^\s*\d+\.\s+", nxt):
                break
            para.append(nxt.strip())
            i += 1
        yield ("paragraph", " ".join(para))


def inline_html(value: str) -> str:
    value = html.escape(value)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    return value


def write_html(blocks):
    out = ["<!doctype html><html lang=\"ru\"><meta charset=\"utf-8\">",
           "<title>G2.GEO — техническое задание для Игоря</title>",
           "<style>body{font:16px/1.55 system-ui,sans-serif;max-width:1060px;margin:36px auto;padding:0 20px;color:#1d2030}h1{font-size:2rem}h2{margin-top:2.2rem;border-bottom:1px solid #dde1ea;padding-bottom:.3rem}h3{margin-top:1.6rem}code,pre{background:#f4f6fa;border-radius:4px}code{padding:.1em .25em}pre{padding:14px;overflow:auto}table{border-collapse:collapse;width:100%;margin:1rem 0}th,td{border:1px solid #cfd5e1;padding:8px;text-align:left;vertical-align:top}th{background:#eef2f8}li{margin:.25rem 0}.meta{color:#5e6677}</style>",
           "<body><p class=\"meta\"><a href=\"../../web/qa-reports/nebula-6-goals-master-plan-2026-06-11/task-tz.html?task=task-g2-geo-content&amp;doc=igor\">Вернуться к карточке G2.GEO</a></p>"]
    for kind, value in blocks:
        if kind != "number":
            active_num_id = None
        if kind == "heading":
            level, text = value
            out.append(f"<h{level}>{inline_html(text)}</h{level}>")
        elif kind == "paragraph":
            out.append(f"<p>{inline_html(value)}</p>")
        elif kind == "bullet":
            out.append(f"<p>• {inline_html(value)}</p>")
        elif kind == "number":
            number, text = value
            out.append(f"<p>{number}. {inline_html(text)}</p>")
        elif kind == "code":
            out.append("<pre>" + html.escape("\n".join(value)) + "</pre>")
        elif kind == "table":
            rows = value
            out.append("<table>")
            for idx, row in enumerate(rows):
                tag = "th" if idx == 0 else "td"
                out.append("<tr>" + "".join(f"<{tag}>{inline_html(cell)}</{tag}>" for cell in row) + "</tr>")
            out.append("</table>")
    out.append("</body></html>\n")
    HTML_OUT.write_text("\n".join(out), encoding="utf-8")


def add_text(paragraph, value: str, *, bold: bool = False, monospace: bool = False):
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Aptos" if not monospace else "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), run.font.name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), run.font.name)
    return run


def write_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    for level in range(1, 4):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(31, 53, 93)
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(cover, "ТЕХНИЧЕСКОЕ ЗАДАНИЕ · ДЛЯ СОГЛАСОВАНИЯ", bold=True)
    numbering = doc.part.numbering_part.element
    base_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(n for n in numbering.num_lst if int(n.numId) == base_num_id)
    abstract_id = base_num.abstractNumId.val
    next_num_id = max(int(n.numId) for n in numbering.num_lst) + 1
    active_num_id = None

    def new_numbering_id():
        nonlocal next_num_id
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(next_num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(abstract_id))
        num.append(abstract)
        numbering.append(num)
        result = next_num_id
        next_num_id += 1
        return result

    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_text(p, text, bold=True)
        elif kind == "paragraph":
            p = doc.add_paragraph()
            add_text(p, value)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_text(p, value)
        elif kind == "number":
            number, text = value
            if number == 1 or active_num_id is None:
                active_num_id = new_numbering_id()
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            numid = OxmlElement("w:numId")
            numid.set(qn("w:val"), str(active_num_id))
            numpr.append(ilvl)
            numpr.append(numid)
            ppr.append(numpr)
            add_text(p, text)
        elif kind == "code":
            p = doc.add_paragraph()
            add_text(p, "\n".join(value), monospace=True)
        elif kind == "table":
            rows = value
            cols = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    target = table.cell(r_idx, c_idx)
                    target.text = ""
                    p = target.paragraphs[0]
                    add_text(p, value, bold=(r_idx == 0))
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    source = SOURCE.read_text(encoding="utf-8")
    parsed = list(parse_blocks(source))
    write_html(parsed)
    write_docx(parsed)
    print(f"generated {HTML_OUT.name} and {DOCX_OUT.name} from {SOURCE.name}")
