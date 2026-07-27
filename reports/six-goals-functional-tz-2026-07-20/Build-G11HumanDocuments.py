from __future__ import annotations

import html
import re
import sys
import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "etalon-tz-g1-1-service-session.md"
HTML_OUT = ROOT / "tz-g1-1-for-igor.html"
DOCX_OUT = ROOT / "tz-g1-1-for-igor.docx"
TASK_ID = "task-g1-session"
TASK_LABEL = "G1.1"
SOURCE_NAME = SOURCE.name
DOCX_NAME = DOCX_OUT.name
DOC_TITLE = "ТЗ G1.1 — Карточка консультации"
DOC_KEYWORDS = "Nebula, G1.1, консультация, ТЗ"
TABLE_HELPER = Path(
    r"C:\Users\alexe\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.723.12215\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_HELPER))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = "123B63"
BLUE = "175CD3"
INK = "17202A"
MUTED = "5F6B7A"
LIGHT_BLUE = "EEF4FF"
TABLE_HEADER = "E8EEF5"
WHITE = "FFFFFF"


def inline_html(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
    return escaped


def metadata_pair(text: str) -> tuple[str, str]:
    label, value = text.split(":", 1)
    label = label.strip().strip("*").strip()
    value = value.strip()
    if value.startswith("**"):
        value = value[2:].lstrip()
    return label, value


def parse_blocks(markdown: str) -> list[dict]:
    lines = markdown.splitlines()
    blocks: list[dict] = []
    i = 0
    seen_section = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            if len(heading.group(1)) == 2:
                seen_section = True
            blocks.append(
                {"type": "heading", "level": len(heading.group(1)), "text": heading.group(2)}
            )
            i += 1
            continue
        if not seen_section and re.match(r"^[^:]{1,48}:\s+.+$", line):
            blocks.append({"type": "paragraph", "text": line.strip()})
            i += 1
            continue
        if line.startswith("> "):
            blocks.append({"type": "quote", "text": line[2:].strip()})
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|(?:\s*:?-+:?\s*\|)+$", lines[i + 1].strip()
        ):
            table_lines = [line]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")] for row in table_lines
            ]
            blocks.append({"type": "table", "rows": rows})
            continue
        if re.match(r"^-\s+", line):
            items = []
            while i < len(lines) and re.match(r"^-\s+", lines[i].strip()):
                items.append(re.sub(r"^-\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "bullets", "items": items})
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "numbers", "items": items})
            continue

        paragraph = [line.strip()]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate:
                break
            if (
                re.match(r"^#{1,3}\s+", candidate)
                or candidate.startswith("> ")
                or candidate.startswith("|")
                or re.match(r"^-\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph.append(candidate)
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraph)})
    return blocks


def slug(text: str) -> str:
    value = re.sub(r"[`«»\"'.,:;!?()/]", "", text.lower())
    value = re.sub(r"[^a-zа-яё0-9]+", "-", value).strip("-")
    return value


def build_html(blocks: list[dict]) -> None:
    title = blocks[0]["text"]
    toc = []
    body = []
    title_seen = False
    metadata_open = False
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not title_seen:
                title_seen = True
                body.append(
                    f'<header class="document-header"><div class="kicker">Nebula · Product specification · {TASK_LABEL}</div>'
                    f"<h1>{inline_html(block['text'])}</h1>"
                    '<p class="lead">Продуктовая постановка для вычитки и согласования с программистом.</p></header>'
                )
                continue
            if metadata_open:
                body.append("</div>")
                metadata_open = False
            anchor = slug(block["text"])
            if level == 2 and block["text"] != "Служебные сведения о документе":
                toc.append((anchor, block["text"]))
            body.append(
                f'<h{level} id="{anchor}">{inline_html(block["text"])}</h{level}>'
            )
            if block["text"] == "Служебные сведения о документе":
                body.append('<div class="metadata metadata-bottom">')
                metadata_open = True
        elif kind == "paragraph":
            text = block["text"]
            if metadata_open and ":" in text:
                label, value = metadata_pair(text)
                body.append(
                    f'<div class="meta-row"><span>{inline_html(label)}</span><b>{inline_html(value)}</b></div>'
                )
            else:
                body.append(f"<p>{inline_html(text)}</p>")
        elif kind in ("bullets", "numbers"):
            tag = "ul" if kind == "bullets" else "ol"
            items = "".join(f"<li>{inline_html(item)}</li>" for item in block["items"])
            body.append(f"<{tag}>{items}</{tag}>")
        elif kind == "quote":
            body.append(f'<aside class="note">{inline_html(block["text"])}</aside>')
        elif kind == "table":
            rows = block["rows"]
            head = "".join(f"<th>{inline_html(cell)}</th>" for cell in rows[0])
            table_body = "".join(
                "<tr>" + "".join(f"<td>{inline_html(cell)}</td>" for cell in row) + "</tr>"
                for row in rows[1:]
            )
            body.append(
                f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{table_body}</tbody></table></div>'
            )
    if metadata_open:
        body.append("</div>")

    toc_html = "".join(
        f'<a href="#{anchor}">{inline_html(text)}</a>' for anchor, text in toc
    )
    page = f"""<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    :root {{ --ink:#17202a; --muted:#5f6b7a; --blue:#175cd3; --navy:#123b63; --line:#dfe5ec; --paper:#fff; --wash:#f5f7fa; }}
    * {{ box-sizing:border-box; }}
    html {{ scroll-behavior:smooth; }}
    body {{ margin:0; color:var(--ink); background:var(--wash); font:17px/1.68 Arial, sans-serif; }}
    .toolbar {{ position:sticky; top:0; z-index:5; display:flex; flex-wrap:wrap; gap:10px; padding:12px 20px; background:rgba(255,255,255,.96); border-bottom:1px solid var(--line); }}
    .toolbar a {{ padding:9px 13px; color:var(--blue); text-decoration:none; font-weight:700; border:1px solid #cad9f5; border-radius:7px; }}
    .toolbar .primary {{ color:#fff; background:var(--blue); border-color:var(--blue); }}
    .layout {{ display:grid; grid-template-columns:260px minmax(0,820px); gap:34px; max-width:1160px; margin:28px auto 60px; padding:0 22px; align-items:start; }}
    nav {{ position:sticky; top:86px; max-height:calc(100vh - 110px); overflow:auto; padding:18px; background:#fff; border:1px solid var(--line); border-radius:12px; }}
    nav b {{ display:block; margin-bottom:10px; color:var(--navy); }}
    nav a {{ display:block; margin:7px 0; color:#445268; text-decoration:none; font-size:14px; line-height:1.35; }}
    nav a:hover {{ color:var(--blue); }}
    article {{ padding:48px 62px 64px; background:var(--paper); border:1px solid var(--line); border-radius:14px; box-shadow:0 8px 30px rgba(28,45,70,.06); }}
    .document-header {{ padding-bottom:24px; margin-bottom:20px; border-bottom:2px solid #d7e3f5; }}
    .kicker {{ color:var(--blue); font-size:13px; font-weight:800; letter-spacing:.08em; text-transform:uppercase; }}
    h1 {{ margin:8px 0 10px; color:var(--navy); font-size:36px; line-height:1.15; }}
    .lead {{ margin:0; color:var(--muted); font-size:18px; }}
    .metadata {{ display:grid; grid-template-columns:1fr 1fr; gap:8px 22px; margin:0 0 30px; padding:18px 20px; background:#f7faff; border:1px solid #dce8f8; border-radius:10px; }}
    .metadata-bottom {{ margin:8px 0 0; }}
    .meta-row {{ display:flex; flex-direction:column; gap:2px; font-size:14px; }}
    .meta-row span {{ color:var(--muted); }}
    h2 {{ margin:42px 0 14px; padding-top:5px; color:var(--navy); font-size:25px; line-height:1.25; }}
    h3 {{ margin:28px 0 10px; color:#244f78; font-size:19px; }}
    p {{ margin:0 0 15px; }}
    ul, ol {{ margin:8px 0 20px; padding-left:28px; }}
    li {{ margin:6px 0; padding-left:4px; }}
    code {{ padding:2px 5px; background:#eef2f6; border-radius:4px; font-size:.9em; }}
    .note {{ margin:20px 0; padding:15px 18px; background:#eef4ff; border-left:4px solid var(--blue); border-radius:6px; }}
    .table-wrap {{ overflow-x:auto; margin:14px 0 24px; }}
    table {{ width:100%; border-collapse:collapse; font-size:15px; }}
    th, td {{ padding:11px 13px; border:1px solid #d7dee8; text-align:left; vertical-align:top; }}
    th {{ color:var(--navy); background:#e8eef5; }}
    @media (max-width:900px) {{ .layout {{ grid-template-columns:1fr; }} nav {{ position:static; max-height:none; }} article {{ padding:30px 24px 45px; }} .metadata {{ grid-template-columns:1fr; }} }}
    @media print {{ body {{ background:#fff; font-size:11pt; }} .toolbar, nav {{ display:none; }} .layout {{ display:block; max-width:none; margin:0; padding:0; }} article {{ border:0; box-shadow:none; padding:0; }} h2 {{ break-after:avoid; }} table {{ break-inside:avoid; }} }}
  </style>
</head>
<body>
  <div class="toolbar">
    <a href="../../web/qa-reports/nebula-6-goals-master-plan-2026-06-11/task-tz.html?task={TASK_ID}&amp;doc=igor">← Карточка {TASK_LABEL}</a>
    <a class="primary" href="{DOCX_NAME}">Скачать DOCX для Игоря</a>
    <a href="{SOURCE_NAME}">Открыть исходник MD</a>
  </div>
  <div class="layout">
    <nav><b>Содержание</b>{toc_html}</nav>
    <article>{''.join(body)}</article>
  </div>
</body>
</html>
"""
    HTML_OUT.write_text(page, encoding="utf-8")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_run(run, size=11, bold=None, color=INK, font="Calibri") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_rich_text(paragraph, text: str, *, size=11, color=INK, bold=False) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run(run, size=size - 0.5, color="31445A", font="Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run(run, size=size, bold=bold, color=color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_run(run, size=9, color=MUTED)


def create_numbering(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    lvl.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = Pt(15)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = Pt(15)


def populate_footer(footer) -> None:
    footer.is_linked_to_previous = False

    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.paragraph_format.space_after = Pt(0)
    set_run(footer_paragraph.add_run(f"{TASK_LABEL}  •  стр. "), size=9, color=MUTED)
    add_page_field(footer_paragraph)


def build_docx(blocks: list[dict]) -> None:
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    configure_styles(doc)

    populate_footer(section.footer)

    title_seen = False
    metadata_mode = False
    for block_index, block in enumerate(blocks):
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not title_seen:
                title_seen = True
                kicker = doc.add_paragraph()
                kicker.paragraph_format.space_before = Pt(7)
                kicker.paragraph_format.space_after = Pt(5)
                set_run(kicker.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ · ДЛЯ СОГЛАСОВАНИЯ"), size=9, bold=True, color=BLUE)
                title = doc.add_paragraph()
                title.paragraph_format.space_after = Pt(5)
                title.paragraph_format.keep_with_next = True
                add_rich_text(title, block["text"], size=24, color=NAVY, bold=True)
                subtitle = doc.add_paragraph()
                subtitle.paragraph_format.space_after = Pt(18)
                set_run(
                    subtitle.add_run("Продуктовая постановка для вычитки и согласования с программистом"),
                    size=12.5,
                    color=MUTED,
                )
                metadata_mode = False
                continue
            metadata_mode = block["text"] == "Служебные сведения о документе"
            style = {1: "Heading 1", 2: "Heading 1", 3: "Heading 2"}[level]
            doc.add_paragraph(block["text"], style=style)
        elif kind == "paragraph":
            text = block["text"]
            if metadata_mode and ":" in text:
                label, value = metadata_pair(text)
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                set_run(p.add_run(f"{label}: "), size=10, bold=True, color=NAVY)
                add_rich_text(p, value, size=10, color=MUTED)
            else:
                metadata_mode = False
                p = doc.add_paragraph()
                if block_index + 1 < len(blocks) and blocks[block_index + 1]["type"] in (
                    "bullets",
                    "numbers",
                ):
                    p.paragraph_format.keep_with_next = True
                add_rich_text(p, text)
        elif kind in ("bullets", "numbers"):
            metadata_mode = False
            for item_index, item in enumerate(block["items"], start=1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.42)
                p.paragraph_format.first_line_indent = Inches(-0.26)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = Pt(15)
                prefix = "•  " if kind == "bullets" else f"{item_index}.  "
                set_run(p.add_run(prefix), size=11, color=INK)
                add_rich_text(p, item)
        elif kind == "quote":
            metadata_mode = False
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_fill(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_rich_text(p, block["text"], color=NAVY)
            apply_table_geometry(
                table,
                [9360],
                indent_dxa=140,
                cell_margins_dxa={"top": 120, "bottom": 120, "start": 160, "end": 160},
            )
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "table":
            metadata_mode = False
            rows = block["rows"]
            cols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for row_idx, row in enumerate(rows):
                for col_idx, value in enumerate(row):
                    cell = table.cell(row_idx, col_idx)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if row_idx == 0:
                        set_cell_fill(cell, TABLE_HEADER)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    add_rich_text(
                        p,
                        value,
                        size=10,
                        color=NAVY if row_idx == 0 else INK,
                        bold=row_idx == 0,
                    )
            widths = [6500, 2860] if cols == 2 else [9360 // cols] * cols
            widths[-1] += 9360 - sum(widths)
            apply_table_geometry(
                table,
                widths,
                indent_dxa=140,
                cell_margins_dxa={"top": 100, "bottom": 100, "start": 140, "end": 140},
            )
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    props = doc.core_properties
    props.title = DOC_TITLE
    props.subject = "Продуктовая постановка для Игоря"
    props.author = "Nebula Product Owner / Project Manager"
    props.keywords = DOC_KEYWORDS
    doc.save(DOCX_OUT)


def main() -> None:
    global SOURCE, HTML_OUT, DOCX_OUT, TASK_ID, TASK_LABEL
    global SOURCE_NAME, DOCX_NAME, DOC_TITLE, DOC_KEYWORDS
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default=SOURCE.name)
    parser.add_argument("--html", default=HTML_OUT.name)
    parser.add_argument("--docx", default=DOCX_OUT.name)
    parser.add_argument("--task-id", default=TASK_ID)
    parser.add_argument("--task-label", default=TASK_LABEL)
    parser.add_argument("--doc-title", default=DOC_TITLE)
    parser.add_argument("--keywords", default=DOC_KEYWORDS)
    args = parser.parse_args()
    SOURCE = ROOT / args.source
    HTML_OUT = ROOT / args.html
    DOCX_OUT = ROOT / args.docx
    TASK_ID = args.task_id
    TASK_LABEL = args.task_label
    SOURCE_NAME = SOURCE.name
    DOCX_NAME = DOCX_OUT.name
    DOC_TITLE = args.doc_title
    DOC_KEYWORDS = args.keywords
    source = SOURCE.read_text(encoding="utf-8")
    blocks = parse_blocks(source)
    build_html(blocks)
    build_docx(blocks)
    print(f"html={HTML_OUT}")
    print(f"docx={DOCX_OUT}")


if __name__ == "__main__":
    main()
